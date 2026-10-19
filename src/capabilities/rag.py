"""Lightweight local document ingestion and retrieval for study prompts."""

from collections import Counter
from dataclasses import dataclass
import math
from pathlib import Path
import re
from typing import Iterable


SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
MAX_FILES = 8
MAX_FILE_SIZE_BYTES = 20 * 1024 * 1024
CHUNK_SIZE = 1_500
CHUNK_OVERLAP = 180
MAX_CONTEXT_CHARS = 12_000


class DocumentProcessingError(ValueError):
   """Raised when an uploaded study document cannot be used."""


@dataclass(frozen=True)
class DocumentChunk:
   """A searchable piece of text with its source attribution."""

   source: str
   content: str
   page: int | None = None


def _tokens(text: str) -> list[str]:
   return re.findall(r"[^\W_]{2,}", text.casefold(), flags=re.UNICODE)


def _split_text(text: str, source: str, page: int | None = None) -> list[DocumentChunk]:
   normalized = re.sub(r"[ \t]+", " ", text).strip()
   if not normalized:
      return []

   chunks = []
   start = 0
   while start < len(normalized):
      end = min(start + CHUNK_SIZE, len(normalized))
      if end < len(normalized):
         boundary = max(
            normalized.rfind("\n", start, end),
            normalized.rfind(". ", start, end),
            normalized.rfind("? ", start, end),
            normalized.rfind("! ", start, end),
            normalized.rfind(" ", start, end),
         )
         if boundary > start + CHUNK_SIZE // 2:
            end = boundary + 1

      content = normalized[start:end].strip()
      if content:
         chunks.append(DocumentChunk(source=source, content=content, page=page))
      if end >= len(normalized):
         break
      start = max(start + 1, end - CHUNK_OVERLAP)
   return chunks


def _extract_text(path: Path) -> list[tuple[int | None, str]]:
   extension = path.suffix.casefold()
   if extension in {".txt", ".md"}:
      return [(None, path.read_text(encoding="utf-8-sig", errors="replace"))]

   if extension == ".pdf":
      try:
         from pypdf import PdfReader
      except ImportError as err:
         raise DocumentProcessingError(
            "PDF support is unavailable. Install the pypdf package and restart the app."
         ) from err
      try:
         reader = PdfReader(str(path))
         return [(number, page.extract_text() or "") for number, page in enumerate(reader.pages, 1)]
      except Exception as err:
         raise DocumentProcessingError(f"Could not read '{path.name}' as a PDF: {err}") from err

   if extension == ".docx":
      try:
         from docx import Document
      except ImportError as err:
         raise DocumentProcessingError(
            "Word document support is unavailable. Install the python-docx package and restart the app."
         ) from err
      try:
         document = Document(str(path))
         sections = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
         for table in document.tables:
            sections.extend(
               " | ".join(cell.text.strip() for cell in row.cells)
               for row in table.rows
            )
         return [(None, "\n".join(sections))]
      except Exception as err:
         raise DocumentProcessingError(
            f"Could not read '{path.name}' as a Word document: {err}"
         ) from err

   raise DocumentProcessingError(
      f"'{path.name}' has an unsupported file type. Upload TXT, Markdown, PDF, or DOCX files."
   )


class DocumentLibrary:
   """In-memory searchable index built from a student's uploaded materials."""

   def __init__(self, chunks: list[DocumentChunk]):
      if not chunks:
         raise DocumentProcessingError(
            "No readable text was found in the uploaded files. Scanned PDFs need OCR before they can be used."
         )
      self.chunks = chunks
      self.document_names = list(dict.fromkeys(chunk.source for chunk in chunks))
      self._chunk_terms = [Counter(_tokens(chunk.content)) for chunk in chunks]
      self._document_frequency = Counter(
         term for terms in self._chunk_terms for term in terms
      )
      self._average_length = (
         sum(sum(terms.values()) for terms in self._chunk_terms) / len(self._chunk_terms)
      )

   @classmethod
   def from_files(cls, files: Iterable[str | Path]) -> "DocumentLibrary":
      if isinstance(files, (str, Path)):
         files = [files]
      paths = [Path(file_path) for file_path in files if file_path]
      if not paths:
         raise DocumentProcessingError("Choose at least one document to upload.")
      if len(paths) > MAX_FILES:
         raise DocumentProcessingError(f"Upload no more than {MAX_FILES} files at a time.")

      chunks: list[DocumentChunk] = []
      for path in paths:
         if not path.is_file():
            raise DocumentProcessingError(f"The uploaded file '{path.name}' is no longer available.")
         if path.suffix.casefold() not in SUPPORTED_EXTENSIONS:
            raise DocumentProcessingError(
               f"'{path.name}' has an unsupported file type. Upload TXT, Markdown, PDF, or DOCX files."
            )
         if path.stat().st_size > MAX_FILE_SIZE_BYTES:
            raise DocumentProcessingError(f"'{path.name}' is larger than the 20 MB per-file limit.")

         try:
            extracted_pages = _extract_text(path)
         except DocumentProcessingError:
            raise
         except OSError as err:
            raise DocumentProcessingError(f"Could not read '{path.name}': {err}") from err

         for page, text in extracted_pages:
            chunks.extend(_split_text(text, path.name, page))

      return cls(chunks)

   def retrieve(self, query: str, max_chunks: int = 5) -> str:
      """Return the most relevant chunks, with source labels, for an AI prompt."""
      query_terms = Counter(_tokens(query))
      if not query_terms:
         selected = list(enumerate(self.chunks[:max_chunks]))
      else:
         scored_chunks = []
         total_chunks = len(self.chunks)
         for index, terms in enumerate(self._chunk_terms):
            score = 0.0
            chunk_length = max(1, sum(terms.values()))
            for term, query_frequency in query_terms.items():
               term_frequency = terms.get(term, 0)
               if not term_frequency:
                  continue
               document_frequency = self._document_frequency[term]
               inverse_frequency = math.log(
                  1 + (total_chunks - document_frequency + 0.5) / (document_frequency + 0.5)
               )
               denominator = term_frequency + 1.2 * (
                  0.25 + 0.75 * chunk_length / max(self._average_length, 1)
               )
               score += inverse_frequency * (term_frequency * 2.2 / denominator) * query_frequency
            scored_chunks.append((score, index, self.chunks[index]))

         relevant = [item for item in scored_chunks if item[0] > 0]
         if relevant:
            best = sorted(relevant, key=lambda item: (-item[0], item[1]))[:max_chunks]
            selected = [(index, chunk) for _, index, chunk in best]
         else:
            selected = list(enumerate(self.chunks[:max_chunks]))

      sections = []
      used_chars = 0
      for _, chunk in selected:
         source = chunk.source
         if chunk.page is not None:
            source += f", page {chunk.page}"
         section = f"[Source: {source}]\n{chunk.content}"
         remaining = MAX_CONTEXT_CHARS - used_chars
         if remaining <= 0:
            break
         sections.append(section[:remaining])
         used_chars += min(len(section), remaining)
      return "\n\n".join(sections)
